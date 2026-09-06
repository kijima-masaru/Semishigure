from io import BytesIO

from openpyxl import load_workbook

from semishigure.report.xlsx import build_workbook, generic_rows


def _run(i: int) -> dict:
    return {
        "id": i,
        "name": f"run{i}",
        "scenario_name": "sc",
        "pbx": {"host": "127.0.0.1", "domain": "pbx.test", "environment": "dev"},
        "started_at": 1700000000.0,
        "finished_at": 1700000100.0,
        "summary": {
            "calls_started": 10 * i, "calls_established": 10 * i, "calls_failed": 0, "fail_by_status": {},
            "invite_to_200_ms": {"p50": 60, "p95": 65, "max": 70}, "answered_by_ext": {"9001": 5}, "answerer_busy_rejects": 0,
            "rtp": {"late_max_ms": 2.5, "late_over_5ms": 0, "tx_packets": 100, "rx_packets": 100, "rx_lost": 0},
            "controller": {"target": 5, "max_concurrency": 50, "ramp_rate": 1, "call_duration": 60, "backoff_reason": ""},
            "monitor": {"threads_by_name": {"freeswitch": {"count": 30, "cpu": 4.0}}, "esl_events": {"CHANNEL_ANSWER": 10}, "hangup_causes": {"NORMAL_CLEARING": 10}, "log_level_counts": {"WARNING": 3}},
            "plugin_rows": [["log_patterns.warnings", 3], ["status_command.x（最終 / 最大）", "1 / 2"]] if i == 1 else [],
        },
        "samples": [
            {"t": 1, "target": 5, "established": 2, "pending": 0, "failed": 0, "cps": 1, "rtp_late_max_ms": 1.0},
            {"t": 2, "target": 5, "established": 5, "pending": 0, "failed": 0, "cps": 0, "rtp_late_max_ms": 2.5},
            {"kind": "monitor", "t": 1, "channels": 4, "cpu": 3.0, "nlwp": 20, "custom": {"a": 1}},
            {"kind": "monitor", "t": 2, "channels": 10, "cpu": 8.0, "nlwp": 30, "custom": {"a": 2}},
        ],
        "calls": [{"id": "c1", "role": "caller", "local": "9100", "remote": "8001", "state": "DONE", "end_reason": "drain", "final_status": 200, "codec": "PCMU", "invite_to_200_ms": 60, "duration_s": 10, "media": {"tx_packets": 500, "rx_packets": 500, "rx_lost": 0, "tx_late_max_ms": 1.2, "tx_late_over_5ms": 0}}],
        "events": [{"t": 0.1, "wall": 1700000000.1, "kind": "start", "message": "controller started"}],
    }


def test_generic_rows_and_workbook():
    rows = dict(generic_rows(_run(1)))
    assert rows["発信 / 確立 / 失敗"] == "10 / 10 / 0"
    assert rows["show channels count"] == "10.0 / 10"  # stable window = later half
    assert rows["PBX プロセス %CPU（区間値）"] == "8.0 / 8.0"
    assert rows["log_patterns.warnings"] == 3
    wb = build_workbook([_run(1), _run(2)])
    buf = BytesIO()
    wb.save(buf)
    wb2 = load_workbook(BytesIO(buf.getvalue()))
    assert wb2.sheetnames == ["記録表", "時系列", "通話一覧", "イベント"]
    ws = wb2["記録表"]
    assert ws["A1"].value == "指標" and ws["B1"].value == "#1 run1" and ws["C1"].value == "#2 run2" and ws["D1"].value == "備考"
    labels = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]
    assert "発信 / 確立 / 失敗" in labels and "log_patterns.warnings" in labels and "判定" in labels
    r = labels.index("発信 / 確立 / 失敗") + 2
    assert ws.cell(row=r, column=2).value == "10 / 10 / 0" and ws.cell(row=r, column=3).value == "20 / 20 / 0"
    r = labels.index("log_patterns.warnings") + 2
    assert ws.cell(row=r, column=2).value == 3 and ws.cell(row=r, column=3).value in ("", None)
    assert labels.index("log_patterns.warnings") < labels.index("判定")
    assert wb2["時系列"].max_row == 5 and wb2["時系列"]["J3"].value == 10  # channels joined by t
    assert wb2["通話一覧"].max_row == 3 and wb2["イベント"].max_row == 3
