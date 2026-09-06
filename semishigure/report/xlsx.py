"""Record sheet export (design §3.6): one column per run, the same row layout as the
manual procedure's 記録表 sheet. Generic rows come from the run statistics and the
PBX monitor; service-specific rows are whatever the plugins reported.

Sheets: 記録表 (comparison), 時系列 (1 s samples of every run), 通話一覧, イベント.
"""

from __future__ import annotations

import time
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F2933")
SECTION_FILL = PatternFill("solid", fgColor="E5E7EB")
THIN = Side(style="thin", color="BFC5CC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _fmt_time(ts: float | None) -> str:
    return time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(ts)) if ts else "-"


def _pair(a: Any, b: Any) -> str:
    return f"{'-' if a is None else a} / {'-' if b is None else b}"


def _series_of(run: dict, kind: str | None) -> list[dict]:
    out = []
    for s in run.get("samples") or []:
        if (s.get("kind") == "monitor") == (kind == "monitor"):
            out.append(s)
    return out


def _stable_window(points: list[dict], key: str, frac: float = 0.5):
    """Values in the later part of the run (after ramp-up), excluding None."""
    vals = [p.get(key) for p in points if p.get(key) is not None]
    if not vals:
        return None, None
    tail = vals[max(0, int(len(vals) * (1 - frac))) :] or vals
    return round(sum(tail) / len(tail), 1), max(vals)


def generic_rows(run: dict) -> list[tuple[str, Any] | tuple[str, None]]:
    """(label, value) rows; a (label, None) tuple with label starting with '§' is a section header."""
    summary = run.get("summary") or {}
    ctl = summary.get("controller") or {}
    load = _series_of(run, None)
    mon = _series_of(run, "monitor")
    est_mean, est_max = _stable_window(load, "established")
    ch_mean, ch_max = _stable_window(mon, "channels")
    cpu_mean, cpu_max = _stable_window(mon, "cpu")
    nlwp_mean, nlwp_max = _stable_window(mon, "nlwp")
    late_max = max((p.get("rtp_late_max_ms") or 0 for p in load), default=None)
    pbx = run.get("pbx") or {}
    monitor = summary.get("monitor") or {}
    rtp = summary.get("rtp") or {}
    i200 = summary.get("invite_to_200_ms") or {}
    rows: list[tuple[str, Any]] = [
        ("§実施情報", None),
        ("実施日 / ラン", f"{_fmt_time(run.get('started_at'))} / {run.get('name')}"),
        ("開始時刻 → 終了時刻", f"{_fmt_time(run.get('started_at'))} → {_fmt_time(run.get('finished_at'))}"),
        ("PBX / 環境", f"{pbx.get('host', '-')} ({pbx.get('domain', '-')}) {pbx.get('environment', '')}"),
        ("シナリオ", run.get("scenario_name")),
        ("目標同時数（最終） / 上限", _pair(ctl.get("target"), ctl.get("max_concurrency"))),
        ("発信レート / 通話長", f"{ctl.get('ramp_rate', '-')}/s / {ctl.get('call_duration', '-')}s"),
        ("§負荷（アプリ側）", None),
        ("発信 / 確立 / 失敗", f"{summary.get('calls_started', '-')} / {summary.get('calls_established', '-')} / {summary.get('calls_failed', '-')}"),
        ("失敗の内訳", ", ".join(f"{k}:{v}" for k, v in (summary.get("fail_by_status") or {}).items()) or "-"),
        ("同時確立数（安定時平均 / 最大）", _pair(est_mean, est_max)),
        ("INVITE→200 p50 / p95 / max (ms)", f"{i200.get('p50', '-')} / {i200.get('p95', '-')} / {i200.get('max', '-')}"),
        ("応答内線の分布", ", ".join(f"{k}:{v}" for k, v in (summary.get("answered_by_ext") or {}).items()) or "-"),
        ("応答側 486（内線上限）", summary.get("answerer_busy_rejects", "-")),
        ("RTP 送出遅れ最大 (ms) / 5ms 超", _pair(rtp.get("late_max_ms", late_max), rtp.get("late_over_5ms"))),
        ("RTP 送受信 / 欠落", f"tx {rtp.get('tx_packets', '-')} / rx {rtp.get('rx_packets', '-')} / lost {rtp.get('rx_lost', '-')}"),
        ("自動減少の理由", ctl.get("backoff_reason") or "-"),
        ("§PBX ホスト（監視、安定時平均 / 最大）", None),
        ("show channels count", _pair(ch_mean, ch_max)),
        ("PBX プロセス %CPU（区間値）", _pair(cpu_mean, cpu_max)),
        ("PBX プロセス nlwp", _pair(nlwp_mean, nlwp_max)),
        ("スレッド別 %CPU（最終）", ", ".join(f"{k}:{v.get('cpu')}({v.get('count')})" for k, v in (monitor.get("threads_by_name") or {}).items()) or "-"),
        ("ESL / AMI イベント", ", ".join(f"{k.replace('CHANNEL_', '')}:{v}" for k, v in (monitor.get("esl_events") or {}).items()) or "-"),
        ("PBX 側の切断理由", ", ".join(f"{k}:{v}" for k, v in (monitor.get("hangup_causes") or {}).items()) or "-"),
        ("ログ WARNING / ERR / CRIT", "{} / {} / {}".format(*[(monitor.get("log_level_counts") or {}).get(k, 0) for k in ("WARNING", "ERR", "CRIT")])),
    ]
    plugin_rows = summary.get("plugin_rows") or []
    if plugin_rows:
        rows.append(("§プラグイン", None))
        rows.extend((str(label), value) for label, value in plugin_rows)
    rows += [("§判定", None), ("頭打ちの兆候（どの指標が最初に劣化したか）", ""), ("この回の有効 / 無効", ""), ("備考", run.get("notes") or "")]
    return rows


def build_workbook(runs: list[dict]) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "記録表"
    # union of row labels, kept inside their sections (plugin rows differ per run)
    per_run = [generic_rows(r) for r in runs]
    by_section: dict[str, list[str]] = {}
    section_order: list[str] = []
    for rows in per_run:
        current = ""
        for label, _ in rows:
            if label.startswith("§"):
                current = label
                if current not in section_order:
                    section_order.append(current)
                by_section.setdefault(current, [])
                continue
            if label not in by_section.setdefault(current, []):
                by_section[current].append(label)
    # 判定 always comes last
    section_order = [x for x in section_order if x != "§判定"] + (["§判定"] if "§判定" in section_order else [])
    labels: list[str] = []
    for sec in section_order:
        labels.append(sec)
        labels.extend(by_section.get(sec, []))
    ws.cell(row=1, column=1, value="指標").font = Font(bold=True, color="FFFFFF")
    ws.cell(row=1, column=1).fill = HEADER_FILL
    for j, run in enumerate(runs, start=2):
        c = ws.cell(row=1, column=j, value=f"#{run.get('id')} {run.get('name')}")
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(row=1, column=len(runs) + 2, value="備考").font = Font(bold=True, color="FFFFFF")
    ws.cell(row=1, column=len(runs) + 2).fill = HEADER_FILL
    maps = [dict(rows) for rows in per_run]
    for i, label in enumerate(labels, start=2):
        c = ws.cell(row=i, column=1, value=label.lstrip("§"))
        if label.startswith("§"):
            c.font = Font(bold=True)
            for j in range(1, len(runs) + 3):
                ws.cell(row=i, column=j).fill = SECTION_FILL
            continue
        for j, m in enumerate(maps, start=2):
            v = m.get(label, "")
            ws.cell(row=i, column=j, value=v if isinstance(v, int | float | str) or v is None else str(v))
    for row in ws.iter_rows(min_row=1, max_row=len(labels) + 1, max_col=len(runs) + 2):
        for c in row:
            c.border = BORDER
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 44
    for j in range(2, len(runs) + 3):
        ws.column_dimensions[get_column_letter(j)].width = 30
    ws.freeze_panes = "B2"

    _sheet_series(wb, runs)
    _sheet_calls(wb, runs)
    _sheet_events(wb, runs)
    return wb


def _sheet_series(wb: Workbook, runs: list[dict]) -> None:
    ws = wb.create_sheet("時系列")
    head = ["run", "t(s)", "target", "established", "pending", "failed", "calls/s", "rtp_late_max_ms", "pump_late_max_ms", "channels", "pbx_cpu", "pbx_nlwp", "custom"]
    ws.append(head)
    for run in runs:
        rid = f"#{run.get('id')}"
        mon = {round(p["t"]): p for p in _series_of(run, "monitor")}
        for p in _series_of(run, None):
            m = mon.get(round(p["t"])) or {}
            custom = m.get("custom") or {}
            ws.append([rid, p.get("t"), p.get("target"), p.get("established"), p.get("pending"), p.get("failed"), p.get("cps"), p.get("rtp_late_max_ms"), p.get("pump_late_max_ms"), m.get("channels"), m.get("cpu"), m.get("nlwp"), ", ".join(f"{k}={v}" for k, v in custom.items()) if custom else ""])
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "A2"


def _sheet_calls(wb: Workbook, runs: list[dict]) -> None:
    ws = wb.create_sheet("通話一覧")
    head = ["run", "id", "role", "local", "remote", "extension", "state", "end_reason", "status", "codec", "INVITE→100 ms", "INVITE→180 ms", "INVITE→200 ms", "200→RTP tx ms", "200→RTP rx ms", "duration s", "tx pkts", "rx pkts", "rx lost", "tx late max ms", "tx late >5ms"]
    ws.append(head)
    for run in runs:
        for c in run.get("calls") or []:
            m = c.get("media") or {}
            ws.append([f"#{run.get('id')}", c.get("id"), c.get("role"), c.get("local"), c.get("remote"), c.get("extension"), c.get("state"), c.get("end_reason"), c.get("final_status"), c.get("codec"), c.get("invite_to_100_ms"), c.get("invite_to_180_ms"), c.get("invite_to_200_ms"), c.get("established_to_first_rtp_tx_ms"), c.get("established_to_first_rtp_rx_ms"), c.get("duration_s"), m.get("tx_packets"), m.get("rx_packets"), m.get("rx_lost"), m.get("tx_late_max_ms"), m.get("tx_late_over_5ms")])
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "A2"


def _sheet_events(wb: Workbook, runs: list[dict]) -> None:
    ws = wb.create_sheet("イベント")
    ws.append(["run", "t(s)", "time", "kind", "message"])
    for run in runs:
        for e in run.get("events") or []:
            ws.append([f"#{run.get('id')}", e.get("t"), _fmt_time(e.get("wall")), e.get("kind"), e.get("message")])
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "A2"
